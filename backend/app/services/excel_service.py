import io
import re
import os
from typing import List, Dict, Tuple, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pandas as pd

def normalize_indian_mobile(raw_val: Any) -> Tuple[bool, str, str]:
    """
    Normalizes Indian mobile numbers to canonical format: +91XXXXXXXXXX.
    Returns: (is_valid, normalized_number, reason_if_invalid)
    """
    if raw_val is None:
        return False, "", "Empty value"
    
    val_str = str(raw_val).strip()
    if not val_str or val_str.lower() in ["nan", "null", "none"]:
        return False, "", "Empty value"
    
    # In Excel, numbers might be read as floats like 9876543210.0
    if val_str.endswith(".0"):
        val_str = val_str[:-2]
        
    # Strip spaces, hyphens, dots, parentheses
    cleaned = re.sub(r"[\s\-\.\(\)]+", "", val_str)
    
    # Check if contains letters or invalid symbols
    digits_only = re.sub(r"^\+", "", cleaned)
    if not digits_only.isdigit():
        return False, val_str, "Contains non-numeric characters"
        
    # Handle prefixes
    if cleaned.startswith("+91"):
        number_part = cleaned[3:]
    elif cleaned.startswith("91") and len(cleaned) == 12:
        number_part = cleaned[2:]
    elif cleaned.startswith("0") and len(cleaned) == 11:
        number_part = cleaned[1:]
    elif len(cleaned) == 10:
        number_part = cleaned
    else:
        return False, val_str, f"Invalid length ({len(cleaned)} digits)"

    # Validate 10-digit Indian mobile number: starts with 6, 7, 8, or 9
    if len(number_part) != 10:
        return False, val_str, f"Expected 10 digits, got {len(number_part)}"
        
    if number_part[0] not in ["6", "7", "8", "9"]:
        return False, val_str, "Indian mobile numbers must start with 6, 7, 8, or 9"
        
    return True, f"+91{number_part}", ""

class ExcelService:
    @staticmethod
    def detect_mobile_column(df: pd.DataFrame) -> str:
        """Detect the column header most likely containing mobile numbers."""
        potential_names = [
            "mobile", "mobile_number", "mobilenumber", "phone", "phonenumber", 
            "phone_number", "contact", "contact_number", "cell", "cellphone", "numbers"
        ]
        
        # Check column names (lowercased, stripped)
        for col in df.columns:
            clean_col = str(col).lower().replace(" ", "_").replace("-", "_")
            if any(name in clean_col for name in potential_names):
                return col
                
        # If no obvious column header matches, inspect first few rows for phone-like digits
        for col in df.columns:
            sample = df[col].dropna().head(5)
            match_count = 0
            for val in sample:
                is_valid, _, _ = normalize_indian_mobile(val)
                if is_valid:
                    match_count += 1
            if match_count >= 2:
                return col
                
        # Fallback to the first column
        return df.columns[0]

    @classmethod
    def parse_and_validate(cls, file_path: str) -> Dict[str, Any]:
        """
        Parses .xlsx or .xls file and validates mobile numbers.
        Returns dict with valid_numbers, stats, and sample invalid items.
        """
        if file_path.endswith(".csv"):
            df = pd.read_csv(file_path, dtype=str)
        else:
            df = pd.read_excel(file_path, dtype=str)
            
        if df.empty:
            return {
                "total_numbers": 0,
                "valid_numbers": 0,
                "invalid_numbers": 0,
                "duplicate_numbers": 0,
                "valid_list": [],
                "sample_valid": [],
                "sample_invalid": []
            }

        target_col = cls.detect_mobile_column(df)
        raw_values = df[target_col].tolist()
        
        total_count = len(raw_values)
        seen_valid = set()
        valid_list: List[str] = []
        duplicate_count = 0
        invalid_list: List[Dict[str, Any]] = []

        for idx, raw in enumerate(raw_values, start=2): # row 2 in Excel
            is_valid, normalized, reason = normalize_indian_mobile(raw)
            if not is_valid:
                invalid_list.append({
                    "row": idx,
                    "value": str(raw) if raw is not None else "",
                    "reason": reason
                })
            else:
                if normalized in seen_valid:
                    duplicate_count += 1
                else:
                    seen_valid.add(normalized)
                    valid_list.append(normalized)

        return {
            "total_numbers": total_count,
            "valid_numbers": len(valid_list),
            "invalid_numbers": len(invalid_list),
            "duplicate_numbers": duplicate_count,
            "valid_list": valid_list,
            "sample_valid": valid_list[:10],
            "sample_invalid": invalid_list[:10]
        }

    @staticmethod
    def generate_sample_file() -> io.BytesIO:
        """Generates a sample Excel template for users to download."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sample Mobile Numbers"

        # Headers
        headers = ["mobile_number", "customer_name", "notes"]
        ws.append(headers)

        # Style header
        header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        sample_rows = [
            ["9876543210", "Aarav Sharma", "Delhi Customer"],
            ["9876543211", "Diya Patel", "Mumbai Customer"],
            ["+919876543212", "Rohan Verma", "Bengaluru Customer"],
            ["919876543213", "Ananya Reddy", "Hyderabad Customer"],
            ["9876543214", "Vikram Malhotra", "Kolkata Customer"],
            ["9876543215", "Pooja Joshi", "Pune Customer"],
            ["9876543216", "Sanjay Singhania", "Ahmedabad Customer"],
            ["9876543217", "Neha Kapoor", "Jaipur Customer"],
            ["9876543218", "Arjun Nair", "Kochi Customer"],
            ["9876543219", "Kavita Rao", "Chennai Customer"]
        ]

        for row in sample_rows:
            ws.append(row)

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 25

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_results_excel(job: Any, results: List[Any], file_path: str):
        """
        Exports job results into a styled .xlsx report file.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Account Check Results"

        # Metadata Header
        ws.append(["Groupin Account Checker — Verification Report"])
        ws.append(["Job ID:", job.job_id])
        ws.append(["Source File:", job.filename])
        ws.append(["Total Checked:", job.processed_numbers])
        ws.append(["Accounts Found:", job.accounts_found])
        ws.append(["Not Registered:", job.not_registered])
        ws.append(["Failed:", job.failed])
        ws.append([])  # empty row

        # Table Header
        headers = [
            "#", "Mobile Number", "Groupin Account", "Groupin User ID", 
            "Name", "Status", "Checked At", "Error Details"
        ]
        ws.append(headers)
        header_row_idx = 9

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Fills for badges
        green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        green_font = Font(name="Calibri", size=10, bold=True, color="166534")
        red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        red_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

        for idx, res in enumerate(results, start=1):
            row_num = header_row_idx + idx
            acct_str = "YES" if res.groupin_account_exists else "NO"
            ws.append([
                idx,
                res.mobile_number,
                acct_str,
                res.groupin_user_id or "-",
                res.name or "-",
                res.status or "-",
                res.checked_at.strftime("%Y-%m-%d %H:%M:%S") if res.checked_at else "-",
                res.error or ""
            ])

            # Apply badge colors on Groupin Account column
            acct_cell = ws.cell(row=row_num, column=3)
            if res.groupin_account_exists:
                acct_cell.fill = green_fill
                acct_cell.font = green_font
            else:
                acct_cell.fill = red_fill
                acct_cell.font = red_font
            acct_cell.alignment = Alignment(horizontal="center")

        # Column widths
        widths = [8, 20, 18, 20, 24, 16, 22, 30]
        for col_idx, width in enumerate(widths, start=1):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        wb.save(file_path)
